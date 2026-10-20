#!/usr/bin/env python3
"""Converte o export do Power BI do campeonato para os arquivos que o app lê.

Quando o sistema de vendas não atualiza, o Power BI do campeonato é a saída. Ele
exporta UM arquivo por dia ('19-10-2026.xlsx'), com uma aba por indicador e as
colunas META | TIME | RESULTADO_ATUAL (um número por loja). O app espera outro
padrão: um arquivo por indicador, uma linha por loja e uma coluna por dia
('19/10/2026 (Seg)'), com o mesmo nome do arquivo da semana anterior.

    python3 importar_powerbi.py                    # acha a rodada sozinho
    python3 importar_powerbi.py --pasta ".../SEMANA ATUAL/rodada 17"
    python3 importar_powerbi.py --modo dia         # se o Power BI vier por dia
    python3 importar_powerbi.py --simular          # só mostra, não grava

Como usar no dia a dia: salve o export de cada dia na pasta da rodada, com a
data no nome (19-10-2026.xlsx, 20-10-2026.xlsx...), e rode o comando. Ele relê
TODOS os arquivos de data da pasta e refaz os arquivos por indicador, então pode
rodar de novo à vontade (é idempotente). Depois é só 'Reprocessar' no site.

Modo de leitura do RESULTADO_ATUAL:
  acumulado (padrão): o valor é o total da semana até aquele dia; o dia vale a
                      diferença para o arquivo anterior.
  dia               : o valor é só o do dia.
Na segunda-feira os dois modos dão o mesmo resultado; confira na terça (ver o
aviso que o script imprime quando o total não cresce como acumulado).
"""
import argparse
import re
import sys
import unicodedata
from datetime import date, datetime, timedelta
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl

DIAS = ["Seg", "Ter", "Qua", "Qui", "Sex", "Sáb", "Dom"]
RE_DATA = re.compile(r"^(\d{2})-(\d{2})-(\d{4})\.xlsx$", re.I)
SIMILARIDADE_MIN = 0.6


def achar_origem():
    base = Path.home() / "Library/CloudStorage"
    for p in base.glob("OneDrive-*/Claude/Campeonato Petz"):
        if p.is_dir():
            return p
    return None


def norm(t):
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]", "", t.upper())


def snapshots(pasta: Path):
    """[(data, arquivo)] dos 'dd-mm-aaaa.xlsx' da pasta, em ordem de data."""
    achados = []
    for f in pasta.glob("*.xlsx"):
        m = RE_DATA.match(f.name)
        if m and not f.name.startswith("~"):
            achados.append((date(int(m[3]), int(m[2]), int(m[1])), f))
    return sorted(achados)


def achar_pasta(origem: Path):
    """A rodada mais recente de SEMANA ATUAL que tem export com data no nome."""
    cands = []
    for p in (origem / "SEMANA ATUAL").glob("rodada *"):
        try:
            n = int(p.name.split()[-1])
        except ValueError:
            continue
        if snapshots(p):
            cands.append((n, p))
    return max(cands)[1] if cands else None


def ler_snapshot(arq: Path):
    """{aba: {loja: valor}} de um export do Power BI."""
    wb = openpyxl.load_workbook(arq, data_only=True, read_only=True)
    saida = {}
    for ws in wb.worksheets:
        linhas = list(ws.iter_rows(values_only=True))
        if not linhas:
            continue
        cab = [norm(c) for c in linhas[0]]
        if "TIME" not in cab or "RESULTADOATUAL" not in cab:
            print(f"   ⚠️  aba '{ws.title}' sem colunas TIME/RESULTADO_ATUAL — ignorada.")
            continue
        it, iv = cab.index("TIME"), cab.index("RESULTADOATUAL")
        vals = {}
        for r in linhas[1:]:
            if not r or len(r) <= max(it, iv) or not r[it]:
                continue
            v = r[iv]
            try:
                vals[str(r[it]).strip()] = None if v in (None, "-", "") else float(v)
            except (TypeError, ValueError):
                vals[str(r[it]).strip()] = None
        saida[ws.title] = vals
    wb.close()
    return saida


def nome_do_arquivo(aba, anteriores):
    """Nome de saída: o do arquivo da semana anterior que casa com a aba."""
    alvo = norm(aba)
    for f in anteriores:
        if norm(f.stem) == alvo:
            return f.name
    melhor, score = None, 0.0
    for f in anteriores:
        s = SequenceMatcher(None, alvo, norm(f.stem)).ratio()
        if s > score:
            melhor, score = f, s
    if melhor is not None and score >= SIMILARIDADE_MIN:
        return melhor.name
    return aba.upper() + ".xlsx"


def diarios(por_data, modo):
    """{aba: {loja: {data: valor_do_dia}}} a partir dos exports de cada data."""
    datas = sorted(por_data)
    abas = {a for d in datas for a in por_data[d]}
    saida = {a: {} for a in abas}
    avisos = []
    for aba in sorted(abas):
        anterior = {}
        total_ant = None
        for d in datas:
            atual = por_data[d].get(aba)
            if atual is None:
                continue
            for loja, v in atual.items():
                dia = v if modo == "dia" or v is None else v - (anterior.get(loja) or 0.0)
                saida[aba].setdefault(loja, {})[d] = None if v is None else round(dia, 6)
            if modo == "acumulado":
                total = sum(v for v in atual.values() if v)
                if total_ant is not None and total < total_ant * 1.05:
                    avisos.append(f"{aba}: em {d:%d/%m} o total ({total:,.0f}) não cresceu sobre o dia "
                                  f"anterior ({total_ant:,.0f}) — o Power BI pode estar vindo POR DIA; "
                                  f"se for o caso, rode com --modo dia.")
                total_ant = total
                anterior = {**anterior, **{l: (v if v is not None else anterior.get(l, 0.0))
                                           for l, v in atual.items()}}
    return saida, avisos


def gravar(destino: Path, lojas, datas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Sigla Loja", "Data Dia Sem"] + [f"{d:%d/%m/%Y} ({DIAS[d.weekday()]})" for d in datas])
    for loja in sorted(lojas):
        linha = [loja, None]
        for d in datas:
            v = lojas[loja].get(d)
            linha.append("-" if v is None else v)
        ws.append(linha)
    for row in ws.iter_rows(min_row=2, min_col=3):
        for c in row:
            c.number_format = "#,##0"
    wb.save(destino)


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--pasta", help="pasta da rodada (SEMANA ATUAL/rodada N)")
    ap.add_argument("--origem", help="pasta do campeonato no OneDrive")
    ap.add_argument("--modo", choices=["acumulado", "dia"], default="acumulado")
    ap.add_argument("--simular", action="store_true", help="não grava nada")
    args = ap.parse_args()

    origem = Path(args.origem) if args.origem else achar_origem()
    pasta = Path(args.pasta) if args.pasta else (achar_pasta(origem) if origem else None)
    if not pasta or not pasta.is_dir():
        print("❌ Não achei a pasta da rodada com arquivos 'dd-mm-aaaa.xlsx'. Use --pasta.")
        return 1
    snaps = snapshots(pasta)
    if not snaps:
        print(f"❌ Nenhum arquivo 'dd-mm-aaaa.xlsx' em {pasta}")
        return 1

    # A semana vai de segunda a domingo; as colunas só vão até o último export.
    segunda = snaps[0][0] - timedelta(days=snaps[0][0].weekday())
    fora = [d for d, _ in snaps if not (segunda <= d < segunda + timedelta(days=7))]
    if fora:
        print(f"❌ Exports de semanas diferentes na mesma pasta: {', '.join(f'{d:%d/%m}' for d in fora)}")
        return 1
    ate = snaps[-1][0]
    datas = [segunda + timedelta(days=i) for i in range((ate - segunda).days + 1)]

    rod = pasta.name.split()[-1]
    anteriores = []
    if origem:
        pa = origem / "SEMANA ANTERIOR" / f"rodada {rod}"
        if pa.is_dir():
            anteriores = [f for f in sorted(pa.glob("*.xlsx")) if not f.name.startswith("~")]
    print(f"📁 Pasta : {pasta}")
    print(f"📅 Semana: {datas[0]:%d/%m} a {datas[-1]:%d/%m} · {len(snaps)} export(s): "
          + ", ".join(f"{d:%d/%m}" for d, _ in snaps))
    print(f"🔎 Modo  : {args.modo}\n")

    por_data = {d: ler_snapshot(f) for d, f in snaps}
    saida, avisos = diarios(por_data, args.modo)

    for aba in sorted(saida):
        nome = nome_do_arquivo(aba, anteriores)
        lojas = saida[aba]
        totais = [sum(v for l in lojas.values() if (v := l.get(d))) for d in datas]
        resumo = " | ".join(f"{d:%d/%m}: R$ {t:,.0f}" for d, t in zip(datas, totais))
        base = " (nome igual ao da semana anterior)" if any(f.name == nome for f in anteriores) \
            else " ⚠️ sem arquivo da semana anterior com esse nome"
        print(f"✔ {aba:32s} → {nome}{base}\n    {len(lojas)} lojas · {resumo}")
        if not args.simular:
            gravar(pasta / nome, lojas, datas)

    for a in avisos:
        print(f"\n⚠️  {a}")
    print("\n" + ("(simulação: nada foi gravado)" if args.simular
                  else "✅ Arquivos gravados na pasta da rodada. Agora é só reprocessar no site."))
    return 0


if __name__ == "__main__":
    sys.exit(main())
